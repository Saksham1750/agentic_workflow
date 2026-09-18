import uuid
from docx import Document as DocxDocument


class DOCXParser:
    def parse(self, content: bytes) -> dict:
        import io
        doc = DocxDocument(io.BytesIO(content))
        sections = []
        current_section = {"id": str(uuid.uuid4()), "title": "Untitled", "page": 0, "content": ""}

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            if para.style and para.style.name.startswith("Heading"):
                if current_section["content"].strip():
                    sections.append(current_section)
                current_section = {
                    "id": str(uuid.uuid4()),
                    "title": text,
                    "page": 0,
                    "content": "",
                }
            else:
                current_section["content"] += text + "\n"

        if current_section["content"].strip():
            sections.append(current_section)

        if not sections:
            full_text = "\n".join(p.text for p in doc.paragraphs)
            sections = [{"id": str(uuid.uuid4()), "title": "Document", "page": 1, "content": full_text}]

        return {"sections": sections, "total_pages": 1}
