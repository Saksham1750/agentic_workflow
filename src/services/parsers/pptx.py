import uuid
from pptx import Presentation


class PPTXParser:
    def parse(self, content: bytes) -> dict:
        import io
        prs = Presentation(io.BytesIO(content))
        sections = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = []
            title = None

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_text.append(text)

                if shape.has_title and shape.title.text.strip():
                    title = shape.title.text.strip()

            if slide_text:
                sections.append({
                    "id": str(uuid.uuid4()),
                    "title": title or f"Slide {slide_num}",
                    "page": slide_num,
                    "content": "\n".join(slide_text),
                })

        if not sections:
            sections = [{"id": str(uuid.uuid4()), "title": "Presentation", "page": 1, "content": ""}]

        return {"sections": sections, "total_pages": len(prs.slides)}
