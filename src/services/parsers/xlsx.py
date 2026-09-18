import uuid
from openpyxl import load_workbook


class XLSXParser:
    def parse(self, content: bytes) -> dict:
        import io
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sections = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_text = [str(cell) if cell is not None else "" for cell in row]
                if any(row_text):
                    rows.append(" | ".join(row_text))

            if rows:
                sections.append({
                    "id": str(uuid.uuid4()),
                    "title": sheet_name,
                    "page": 1,
                    "content": "\n".join(rows),
                })

        wb.close()

        if not sections:
            sections = [{"id": str(uuid.uuid4()), "title": "Spreadsheet", "page": 1, "content": ""}]

        return {"sections": sections, "total_pages": len(wb.sheetnames)}
